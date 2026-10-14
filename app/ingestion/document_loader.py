"""Native multi-format document loader and text parser."""

import csv
import io
from pathlib import Path
from typing import List
from bs4 import BeautifulSoup
import docx
import openpyxl
import pptx
from striprtf.striprtf import rtf_to_text

from app.core.exceptions import OCRError
from app.core.logging import get_logger
from app.ocr.engine_factory import get_ocr_engine
from app.ocr.image_preprocessor import get_image_preprocessor
from app.ocr.pdf_extractor import get_pdf_extractor
from app.ocr.table_extractor import TableExtractor
from app.schemas.ocr import OCRBlock, OCRBoundingBox, OCRPage, OCRTable
from app.storage.artifact_store import get_artifact_store

logger = get_logger("document_loader")


class DocumentLoader:
    """Dispatches extraction according to file format and structure."""

    def __init__(self) -> None:
        self.pdf_extractor = get_pdf_extractor()
        self.preprocessor = get_image_preprocessor()
        self.artifact_store = get_artifact_store()
        self.table_extractor = TableExtractor()

    def load_and_extract(
        self,
        document_id: str,
        filename: str,
        file_bytes: bytes,
        preferred_engine: str = "hybrid"
    ) -> List[OCRPage]:
        """Extract content into structured OCRPage list based on file extension."""
        ext = Path(filename).suffix.lstrip(".").lower()

        # 1. PDF
        if ext == "pdf":
            return self.pdf_extractor.process_pdf_bytes(document_id, file_bytes, preferred_engine)

        # 2. Image formats
        if ext in ["png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"]:
            return self._extract_image(document_id, filename, file_bytes, preferred_engine)

        # 3. Word Document (.docx)
        if ext == "docx":
            return self._extract_docx(document_id, file_bytes)

        # 4. Excel Spreadsheet (.xlsx, .csv)
        if ext in ["xlsx", "csv"]:
            return self._extract_spreadsheet(document_id, file_bytes, ext)

        # 5. PowerPoint Presentation (.pptx)
        if ext == "pptx":
            return self._extract_pptx(file_bytes)

        # 6. RTF Document
        if ext == "rtf":
            return self._extract_rtf(file_bytes)

        # 7. HTML / XML Document
        if ext in ["html", "htm", "xml"]:
            return self._extract_markup(file_bytes)

        # 8. Plain text / Markdown (.txt, .md)
        return self._extract_plaintext(file_bytes)

    def _extract_image(
        self,
        document_id: str,
        filename: str,
        file_bytes: bytes,
        preferred_engine: str
    ) -> List[OCRPage]:
        """Process a standalone image through preprocessing and OCR."""
        cv_img = self.preprocessor.from_bytes(file_bytes)

        # 1. Enhance for vision while strictly preserving color fidelity and correcting skew
        vision_img, warnings, metrics = self.preprocessor.preprocess_for_vision(cv_img)

        # Save processed image to MinIO
        proc_bytes = self.preprocessor.to_bytes(vision_img, ".png")
        image_key = self.artifact_store.save_page_image(document_id, 1, proc_bytes)

        # 2. Run OCR Engine (Hybrid routes color to Gemini Vision, and preprocessed grayscale to Tesseract)
        engine = get_ocr_engine(preferred_engine)
        text, avg_conf, blocks = engine.extract_text_and_blocks(vision_img, 1)

        if avg_conf < 60.0:
            warnings.append("LOW_OCR_CONFIDENCE")

        # 3. Extract structured tables: morphology + Markdown table parsing
        tables = self.table_extractor.extract_from_image(vision_img)
        md_tables = TableExtractor.extract_tables_from_markdown(text, start_index=len(tables))
        for m_tbl in md_tables:
            if not any(t.headers == m_tbl.headers for t in tables):
                tables.append(m_tbl)

        for tbl in tables:
            if tbl.csv_data:
                try:
                    self.artifact_store.save_table_csv(document_id, 1, tbl.table_index, tbl.csv_data)
                except Exception:
                    pass

        return [
            OCRPage(
                page_number=1,
                ocr_engine=engine.engine_name,
                confidence=avg_conf,
                char_count=len(text),
                table_count=len(tables),
                has_selectable_text=False,
                raw_text=text,
                blocks=blocks,
                tables=tables,
                warnings=warnings,
                image_path=image_key
            )
        ]

    def _extract_docx(self, document_id: str, file_bytes: bytes) -> List[OCRPage]:
        """Extract paragraphs and tables natively from Microsoft Word .docx."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)

            tables: List[OCRTable] = []
            for idx, tbl in enumerate(doc.tables):
                rows_data = []
                for row in tbl.rows:
                    rows_data.append([c.text.strip() for c in row.cells])
                if rows_data:
                    headers = rows_data[0]
                    content_rows = rows_data[1:] if len(rows_data) > 1 else []
                    tbl_obj = OCRTable(
                        table_index=idx,
                        num_rows=len(rows_data),
                        num_cols=len(headers),
                        headers=headers,
                        rows=content_rows,
                        csv_data=TableExtractor.format_csv_table(headers, content_rows),
                        markdown_data=TableExtractor.format_markdown_table(headers, content_rows),
                        confidence=100.0
                    )
                    tables.append(tbl_obj)
                    if tbl_obj.csv_data:
                        try:
                            self.artifact_store.save_table_csv(document_id, 1, idx, tbl_obj.csv_data)
                        except Exception:
                            pass

            return [
                OCRPage(
                    page_number=1,
                    ocr_engine="NATIVE_DOCX",
                    confidence=100.0,
                    char_count=len(full_text),
                    table_count=len(tables),
                    has_selectable_text=True,
                    raw_text=full_text,
                    blocks=[],
                    tables=tables,
                    warnings=[]
                )
            ]
        except Exception as exc:
            logger.error(f"Docx extraction error: {exc}")
            raise OCRError(f"Failed parsing DOCX: {exc}")


    def _extract_spreadsheet(self, document_id: str, file_bytes: bytes, ext: str) -> List[OCRPage]:
        """Extract worksheets natively as structured tables and Markdown."""
        tables: List[OCRTable] = []
        combined_text_parts = []

        if ext == "csv":
            text_content = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text_content))
            rows = list(reader)
            if rows:
                headers = rows[0]
                body = rows[1:] if len(rows) > 1 else []
                md = TableExtractor.format_markdown_table(headers, body)
                tbl_obj = OCRTable(
                    table_index=0,
                    num_rows=len(rows),
                    num_cols=len(headers),
                    headers=headers,
                    rows=body,
                    csv_data=text_content,
                    markdown_data=md,
                    confidence=100.0
                )
                tables.append(tbl_obj)
                try:
                    self.artifact_store.save_table_csv(document_id, 1, 0, text_content)
                except Exception:
                    pass
                combined_text_parts.append(md)
        else:
            # Excel (.xlsx)
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(c or "").strip() for c in row]
                    if any(row_vals):
                        rows.append(row_vals)
                if rows:
                    headers = rows[0]
                    body = rows[1:] if len(rows) > 1 else []
                    md = TableExtractor.format_markdown_table(headers, body)
                    csv_str = TableExtractor.format_csv_table(headers, body)
                    tbl_obj = OCRTable(
                        table_index=idx,
                        num_rows=len(rows),
                        num_cols=len(headers),
                        headers=headers,
                        rows=body,
                        csv_data=csv_str,
                        markdown_data=md,
                        confidence=100.0
                    )
                    tables.append(tbl_obj)
                    if csv_str:
                        try:
                            self.artifact_store.save_table_csv(document_id, 1, idx, csv_str)
                        except Exception:
                            pass
                    combined_text_parts.append(f"### Sheet: {sheet_name}\n{md}")

        full_text = "\n\n".join(combined_text_parts)
        return [
            OCRPage(
                page_number=1,
                ocr_engine="NATIVE_SPREADSHEET",
                confidence=100.0,
                char_count=len(full_text),
                table_count=len(tables),
                has_selectable_text=True,
                raw_text=full_text,
                blocks=[],
                tables=tables,
                warnings=[]
            )
        ]

    def _extract_pptx(self, file_bytes: bytes) -> List[OCRPage]:
        """Extract slide titles and shapes natively from PowerPoint .pptx."""
        prs = pptx.Presentation(io.BytesIO(file_bytes))
        pages: List[OCRPage] = []

        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    slide_texts.append(shape.text_frame.text.strip())
            slide_content = "\n".join(t for t in slide_texts if t)

            pages.append(
                OCRPage(
                    page_number=idx + 1,
                    ocr_engine="NATIVE_PPTX",
                    confidence=100.0,
                    char_count=len(slide_content),
                    table_count=0,
                    has_selectable_text=True,
                    raw_text=slide_content,
                    blocks=[],
                    tables=[],
                    warnings=[]
                )
            )
        return pages

    def _extract_rtf(self, file_bytes: bytes) -> List[OCRPage]:
        """Strip formatting tags and extract plain text from RTF."""
        rtf_str = file_bytes.decode("utf-8", errors="replace")
        text = rtf_to_text(rtf_str).strip()
        return [
            OCRPage(
                page_number=1,
                ocr_engine="NATIVE_RTF",
                confidence=100.0,
                char_count=len(text),
                table_count=0,
                has_selectable_text=True,
                raw_text=text,
                blocks=[],
                tables=[],
                warnings=[]
            )
        ]

    def _extract_markup(self, file_bytes: bytes) -> List[OCRPage]:
        """Parse HTML or XML content via BeautifulSoup."""
        soup = BeautifulSoup(file_bytes, "html.parser")
        text = soup.get_text(separator="\n\n").strip()
        return [
            OCRPage(
                page_number=1,
                ocr_engine="NATIVE_MARKUP",
                confidence=100.0,
                char_count=len(text),
                table_count=0,
                has_selectable_text=True,
                raw_text=text,
                blocks=[],
                tables=[],
                warnings=[]
            )
        ]

    def _extract_plaintext(self, file_bytes: bytes) -> List[OCRPage]:
        """Load UTF-8 plain text or Markdown files."""
        text = file_bytes.decode("utf-8", errors="replace").strip()
        return [
            OCRPage(
                page_number=1,
                ocr_engine="NATIVE_TEXT",
                confidence=100.0,
                char_count=len(text),
                table_count=0,
                has_selectable_text=True,
                raw_text=text,
                blocks=[],
                tables=[],
                warnings=[]
            )
        ]


_document_loader_instance = DocumentLoader()


def get_document_loader() -> DocumentLoader:
    """Retrieve singleton DocumentLoader instance."""
    return _document_loader_instance
